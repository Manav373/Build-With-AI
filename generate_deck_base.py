import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]  # blank layout

# Color Palette Definitions
DARK_BG = RGBColor(7, 31, 20)          # #071F14 (Deep Forest Dark Emerald)
CARD_BG = RGBColor(15, 56, 36)         # #0F3824 (Dark Emerald Card Glass)
CARD_BORDER = RGBColor(34, 100, 65)    # #226441 (Subtle Emerald Border)
ACCENT_GREEN = RGBColor(0, 255, 102)   # #00FF66 (Bright Neon Emerald Accent)
ACCENT_MINT = RGBColor(16, 185, 129)   # #10B981 (Mint Emerald)
ACCENT_GOLD = RGBColor(245, 158, 11)   # #F59E0B (Amber Gold)
ACCENT_ORANGE = RGBColor(239, 68, 68)   # #EF4444 (Sunset Alert)
TEXT_WHITE = RGBColor(255, 255, 255)   # #FFFFFF
TEXT_MUTED = RGBColor(209, 213, 219)   # #D1D5DB
TEXT_SUBTITLE = RGBColor(167, 243, 208)# #A7F3D0
CARD_INNER_BG = RGBColor(20, 70, 45)    # #14462D

def add_blank_slide_with_bg():
    slide = prs.slides.add_slide(blank_slide_layout)
    # Background shape covering full slide
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = DARK_BG
    bg.line.fill.background()
    return slide

def add_header(slide, category_text, main_title_text):
    # Category Banner
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_GREEN
    p_cat.font.name = "Calibri"

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.7))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = main_title_text
    p_title.font.size = Pt(26)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    p_title.font.name = "Arial"

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def add_badge(slide, left, top, width, height, text, bg_color=CARD_INNER_BG, text_color=ACCENT_GREEN):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = text
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = text_color
    p.font.name = "Calibri"

print("Helper functions defined successfully.")
